import os
import logging
import psutil
import pdfplumber
import docx
import openpyxl
import re
import yaml
import csv
import gradio as gr
from PIL import Image
import pytesseract
import requests
from bs4 import BeautifulSoup
from duckduckgo_search import DDGS
from urllib.parse import urlparse
import time

from typing import List, Dict, Optional
from langchain_core.documents import Document
from langchain_core.output_parsers import StrOutputParser
from langchain_core.runnables import RunnablePassthrough, RunnableLambda
from langchain_core.prompts import ChatPromptTemplate
from langchain_community.document_loaders import TextLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_ollama.llms import OllamaLLM
from langchain_community.vectorstores import FAISS


def load_config(config_file: str = "config.yaml") -> dict:
    """Load the YAML config."""
    with open(config_file, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


class RAGPipeline:
    def __init__(
        self,
        model_name: str,
        embedding_model_name: str,
        max_memory_gb: float
    ):
        self.setup_logging()
        self.check_system_memory(max_memory_gb)
        
        # LLM
        self.llm = OllamaLLM(model=model_name)
        
        # Embeddings
        self.embeddings = HuggingFaceEmbeddings(
            model_name=embedding_model_name,
            model_kwargs={'device': 'cpu'}
        )
        
        # Chat prompt template that includes chat_history
        self.chat_prompt = ChatPromptTemplate.from_template("""
You are a helpful AI assistant that answers questions based on the provided documents and web search results.

Conversation so far:
{chat_history}

Relevant context from the documents and web search:
{context}

Now answer the user's latest question: {question}

If the context includes web search results, make sure to cite the sources when providing information.
""")

    def setup_logging(self):
        logging.basicConfig(level=logging.INFO)
        self.logger = logging.getLogger(__name__)

    def check_system_memory(self, max_memory_gb: float):
        available_memory = psutil.virtual_memory().available / (1024 ** 3)
        self.logger.info(f"Available system memory: {available_memory:.1f} GB")
        if available_memory < max_memory_gb:
            self.logger.warning("Memory is below recommended threshold.")

    def extract_text_from_csv(self, filepath: str) -> str:
        """
        Reads the CSV, outputs:
        1) First line with comma-separated headers: "column1,column2"
        2) Each row as "column1: Text,column2: Text"
        """
        lines = []
        try:
            with open(filepath, mode="r", encoding="utf-8", errors="ignore") as f:
                csv_reader = csv.reader(f)
                headers = next(csv_reader, None)
                if not headers:
                    self.logger.warning(f"No headers found in CSV '{filepath}'")
                    return ""
                # First line: headers, comma-separated
                lines.append(",".join(headers))

                # Subsequent rows: "header1: val1,header2: val2"
                for row in csv_reader:
                    row_text = []
                    for i in range(min(len(row), len(headers))):
                        row_text.append(f"{headers[i]}: {row[i]}")
                    lines.append(",".join(row_text))
        except Exception as e:
            self.logger.warning(f"Failed to parse CSV '{filepath}' with error: {e}")
            return ""
        return "\n".join(lines)

    def extract_text_from_pdf(self, filepath: str) -> str:
        try:
            text_list = []
            with pdfplumber.open(filepath) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text() or ""
                    text_list.append(page_text)
            return "\n".join(text_list)
        except Exception as e:
            self.logger.warning(f"pdfplumber failed on '{filepath}' with error: {e}")
            return ""

    def extract_text_from_docx(self, filepath: str) -> str:
        doc = docx.Document(filepath)
        return "\n".join(para.text for para in doc.paragraphs)

    def extract_text_from_xlsx(self, filepath: str) -> str:
        """
        Similar logic to CSV:
        1) First row => comma-separated headers
        2) Subsequent rows => "header1: val1,header2: val2"
        Combines data from all sheets one after another.
        """
        wb = openpyxl.load_workbook(filepath, read_only=True)
        lines = []
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            header = None
            for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                row_values = list(row) if row else []
                # Skip empty rows
                if all(cell is None for cell in row_values):
                    continue
                # First row => set headers
                if row_index == 0:
                    header = [str(col) if col else "UNKNOWN_COLUMN" for col in row_values]
                    lines.append(",".join(header))
                else:
                    # If there's no header, we can't format properly
                    if not header:
                        self.logger.warning(f"No header found in sheet {sheet_name} for file {filepath}")
                        break
                    row_text = []
                    for i in range(min(len(header), len(row_values))):
                        row_text.append(f"{header[i]}: {row_values[i]}")
                    lines.append(",".join(row_text))
        return "\n".join(lines)

    def extract_text_from_txt(self, filepath: str) -> str:
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    def extract_text_from_md(self, filepath: str) -> str:
        """Reads Markdown (.md) files as plain text."""
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    def extract_text_from_image(self, filepath: str) -> str:
        """Extract text from image using Tesseract OCR."""
        try:
            image = Image.open(filepath)
            text = pytesseract.image_to_string(image)
            return text
        except Exception as e:
            self.logger.warning(f"OCR failed on image '{filepath}' with error: {e}")
            return ""

    def search_web(self, query: str, max_results: int = 5) -> List[Dict[str, str]]:
        """Search the web using DuckDuckGo and return results with URLs and snippets."""
        try:
            with DDGS() as ddgs:
                print(ddgs)
                results = list(ddgs.text(query, max_results=max_results))
                
                # Debug: Print the structure of the first result
                if results:
                    self.logger.info(f"First result structure: {results[0]}")
                    self.logger.info(f"Available keys in first result: {list(results[0].keys())}")
                
                # Filter out results with empty or invalid URLs
                valid_results = []
                for i, result in enumerate(results):
                    # Try different possible field names for URL
                    url = (result.get('link') or 
                           result.get('url') or 
                           result.get('href') or 
                           result.get('result') or 
                           '').strip()
                    
                    self.logger.info(f"Result {i}: URL field = '{url}'")
                    
                    if url and url.startswith(('http://', 'https://')):
                        valid_results.append(result)
                        self.logger.info(f"Valid URL found: {url}")
                    else:
                        self.logger.warning(f"Skipping result {i} with invalid URL: '{url}'")
                        # Debug: Show all fields for this result
                        self.logger.info(f"Result {i} all fields: {result}")
                
                return valid_results
        except Exception as e:
            self.logger.error(f"Web search failed: {e}")
            return []

    def extract_text_from_url(self, url: str, timeout: int = 10) -> str:
        """Extract text content from a web page."""
        try:
            # Validate URL
            if not url or not url.strip():
                self.logger.warning("Empty URL provided")
                return ""
            
            url = url.strip()
            if not url.startswith(('http://', 'https://')):
                self.logger.warning(f"Invalid URL scheme: {url}")
                return ""
            
            headers = {
                'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36'
            }
            response = requests.get(url, headers=headers, timeout=timeout)
            response.raise_for_status()
            
            soup = BeautifulSoup(response.content, 'html.parser')
            
            # Remove script and style elements
            for script in soup(["script", "style"]):
                script.decompose()
            
            # Get text content
            text = soup.get_text()
            
            # Clean up whitespace
            lines = (line.strip() for line in text.splitlines())
            chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
            text = ' '.join(chunk for chunk in chunks if chunk)
            
            return text[:5000]  # Limit to first 5000 characters to avoid too long documents
        except requests.exceptions.RequestException as e:
            self.logger.warning(f"Request failed for URL '{url}': {e}")
            return ""
        except Exception as e:
            self.logger.warning(f"Failed to extract text from URL '{url}': {e}")
            return ""

    def search_and_extract_web_content(self, query: str, max_results: int = 3) -> List[Document]:
        """Search the web and extract content from the top results."""
        self.logger.info(f"Searching web for: {query}")
        
        # Search the web
        search_results = self.search_web(query, max_results=max_results)
        
        if not search_results:
            self.logger.warning("No valid search results found")
            return []
        
        self.logger.info(f"Found {len(search_results)} valid search results")
        
        documents = []
        for i, result in enumerate(search_results):
            try:
                # Try different possible field names for URL, title, and body
                url = (result.get('link') or 
                       result.get('url') or 
                       result.get('href') or 
                       result.get('result') or 
                       '').strip()
                
                title = (result.get('title') or 
                        result.get('name') or 
                        result.get('heading') or 
                        '').strip()
                
                snippet = (result.get('body') or 
                          result.get('snippet') or 
                          result.get('description') or 
                          result.get('text') or 
                          '').strip()
                
                self.logger.info(f"Processing result {i+1}: {url}")
                
                # Extract full text from the webpage
                full_text = self.extract_text_from_url(url)
                
                if full_text:
                    # Combine title, snippet, and full text
                    combined_text = f"Title: {title}\n\nSnippet: {snippet}\n\nFull Content: {full_text}"
                    
                    doc = Document(
                        page_content=combined_text,
                        metadata={
                            "source": url,
                            "title": title,
                            "search_query": query,
                            "result_index": i
                        }
                    )
                    documents.append(doc)
                    self.logger.info(f"Successfully extracted content from: {url}")
                else:
                    self.logger.warning(f"No content extracted from: {url}")
                
                # Add a small delay to be respectful to websites
                time.sleep(1)
                
            except Exception as e:
                self.logger.warning(f"Failed to process search result {i}: {e}")
                continue
        
        self.logger.info(f"Successfully processed {len(documents)} documents from web search")
        return documents

    def extract_text(self, filepath: str) -> str:
        ext = os.path.splitext(filepath)[1].lower()
        if ext == ".pdf":
            return self.extract_text_from_pdf(filepath)
        elif ext == ".docx":
            return self.extract_text_from_docx(filepath)
        elif ext == ".xlsx":
            return self.extract_text_from_xlsx(filepath)
        elif ext == ".txt":
            return self.extract_text_from_txt(filepath)
        elif ext == ".csv":
            return self.extract_text_from_csv(filepath)
        elif ext == ".md":
            return self.extract_text_from_md(filepath)
        elif ext in [".png", ".jpg", ".jpeg"]:
            return self.extract_text_from_image(filepath)
        return ""

    def gather_documents_from_folder(self, folder_path: str) -> List[Document]:
        if not os.path.exists(folder_path):
            self.logger.warning(f"Folder does not exist: {folder_path}")
            return []

        docs = []
        for root, _, files in os.walk(folder_path):
            for file in files:
                if file.lower().endswith((".pdf", ".docx", ".xlsx", ".txt", ".csv", ".md", ".png", ".jpg", ".jpeg")):
                    path = os.path.join(root, file)
                    text = self.extract_text(path)
                    if text.strip():
                        docs.append(Document(page_content=text, metadata={"source": path}))
        
        if not docs:
            self.logger.warning(f"No recognized files found in folder: {folder_path}")

        text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=500,
            chunk_overlap=50,
            length_function=len,
            add_start_index=True,
        )
        splits = text_splitter.split_documents(docs)
        self.logger.info(f"Created {len(splits)} document chunks for folder: {folder_path}")
        return splits

    def create_vectorstore(self, documents: List[Document]) -> FAISS:
        batch_size = 32
        if not documents:
            self.logger.warning("No documents provided to create_vectorstore. Creating an empty vector store.")
            return FAISS.from_texts([""], self.embeddings)

        vectorstore = FAISS.from_documents(documents[:batch_size], self.embeddings)
        for i in range(batch_size, len(documents), batch_size):
            batch = documents[i:i + batch_size]
            vectorstore.add_documents(batch)
            self.logger.info(f"Processed batch {i // batch_size + 1}")
        return vectorstore
    
    def get_question(self, inputs: dict) -> str:
        return inputs["question"]

    def setup_rag_chain(self, vectorstore: FAISS):
        retriever = vectorstore.as_retriever(
            search_type="similarity",
            search_kwargs={"k": 2, "fetch_k": 3}
        )

        def format_docs(docs):
            return "\n\n".join(doc.page_content for doc in docs)
        
        rag_chain = (
            {
                "chat_history": RunnablePassthrough(),
                "context": RunnableLambda(self.get_question) | retriever | format_docs,
                "question": RunnablePassthrough(),
            }
            | self.chat_prompt
            | self.llm
            | StrOutputParser()
        )
        return rag_chain
    
    def query(self, chain, chain_inputs: dict) -> str:
        memory_usage = psutil.Process(os.getpid()).memory_info().rss / 1024 / 1024
        self.logger.info(f"Memory usage: {memory_usage:.1f} MB")
        return chain.invoke(chain_inputs)

    def should_search_web(self, question: str) -> bool:
        """Determine if a question requires web search based on keywords and patterns."""
        # Keywords that suggest need for current/recent information
        web_keywords = [
            'latest', 'recent', 'current', 'today', 'yesterday', 'this week', 'this month',
            'news', 'update', 'breaking', 'trending', 'popular', 'best', 'top',
            'price', 'cost', 'where to buy', 'how much', 'compare',
            'weather', 'forecast', 'temperature',
            'stock', 'market', 'crypto', 'bitcoin',
            'movie', 'film', 'show', 'series', 'release date',
            'restaurant', 'hotel', 'travel', 'vacation', 'trip'
        ]
        
        # Questions that typically need current information
        web_patterns = [
            r'\bwhat is the (latest|current|recent)\b',
            r'\bwhat are the (latest|current|recent)\b',
            r'\bwhen (was|is|will)\b',
            r'\bwhere (can|could|should)\b',
            r'\bhow much (does|do|is|are)\b',
            r'\bwhat (is|are) the (best|top|worst)\b',
            r'\bcompare\b',
            r'\bvs\b',
            r'\bversus\b'
        ]
        
        question_lower = question.lower()
        
        # Check for web keywords
        for keyword in web_keywords:
            if keyword in question_lower:
                return True
        
        # Check for web patterns
        for pattern in web_patterns:
            if re.search(pattern, question_lower):
                return True
        
        return False


def build_pipeline(config: dict):
    rag = RAGPipeline(
        model_name=config["model_name"],
        embedding_model_name=config["embedding_model_name"],
        max_memory_gb=config["max_memory_gb"]
    )

    all_docs = []
    folder_paths = config.get("folder_paths", [])
    for folder_path in folder_paths:
        doc_splits = rag.gather_documents_from_folder(folder_path)
        all_docs.extend(doc_splits)

    if not all_docs:
        rag.logger.warning("No documents found in any specified folders. The vectorstore will be empty.")

    index_path = config["index_path"]
    if os.path.exists(index_path):
        rag.logger.info("Loading existing FAISS index from disk...")
        vectorstore = FAISS.load_local(index_path, rag.embeddings, allow_dangerous_deserialization=True)
    else:
        rag.logger.info("Creating a new FAISS index and saving it locally...")
        vectorstore = rag.create_vectorstore(all_docs)
        vectorstore.save_local(index_path)

    chain = rag.setup_rag_chain(vectorstore)
    return rag, chain


def build_hybrid_pipeline(config: dict):
    """Build a hybrid pipeline that can search both local documents and the web."""
    rag = RAGPipeline(
        model_name=config["model_name"],
        embedding_model_name=config["embedding_model_name"],
        max_memory_gb=config["max_memory_gb"]
    )

    # Build local document vectorstore
    all_docs = []
    folder_paths = config.get("folder_paths", [])
    for folder_path in folder_paths:
        doc_splits = rag.gather_documents_from_folder(folder_path)
        all_docs.extend(doc_splits)

    if not all_docs:
        rag.logger.warning("No documents found in any specified folders. The vectorstore will be empty.")

    index_path = config["index_path"]
    if os.path.exists(index_path):
        rag.logger.info("Loading existing FAISS index from disk...")
        vectorstore = FAISS.load_local(index_path, rag.embeddings, allow_dangerous_deserialization=True)
    else:
        rag.logger.info("Creating a new FAISS index and saving it locally...")
        vectorstore = rag.create_vectorstore(all_docs)
        vectorstore.save_local(index_path)

    chain = rag.setup_rag_chain(vectorstore)
    return rag, chain, vectorstore


def main():
    config = load_config("config.yaml")
    rag, chain, vectorstore = build_hybrid_pipeline(config)
    remove_think_flag = config.get("remove_think", True)
    enable_web_search = config.get("enable_web_search", True)

    with gr.Blocks() as demo:
        gr.Markdown("## RAG Chat with Conversation History and Web Search")

        with gr.Row():
            with gr.Column(scale=3):
                chatbot = gr.Chatbot(label="Conversation", type="messages")
                msg = gr.Textbox(label="Your question")
                
                with gr.Row():
                    submit_btn = gr.Button("Send", variant="primary")
                    web_search_btn = gr.Button("🔍 Web Search", variant="secondary")
                    clear_btn = gr.Button("Clear Chat")
                
                status = gr.Textbox(label="Status", interactive=False, value="Ready")
            
            with gr.Column(scale=1):
                gr.Markdown("### Settings")
                web_search_toggle = gr.Checkbox(
                    label="Enable Web Search", 
                    value=enable_web_search,
                    info="Automatically search the web for current/recent information"
                )
                
                gr.Markdown("### Web Search Options")
                gr.Markdown("""
                **Automatic Web Search** is triggered for questions about:
                - Latest/current/recent information
                - News and updates
                - Prices and comparisons
                - Weather and forecasts
                - Movies, shows, release dates
                - Restaurants, hotels, travel
                - Stock market and crypto
                
                **Manual Web Search** button forces web search for any query, regardless of automatic detection.
                """)

        state = gr.State([])

        def chat_with_status(user_message, history, web_search_enabled, manual_web_search=False):
            # Determine if we should search the web
            should_search = (web_search_enabled and rag.should_search_web(user_message)) or manual_web_search
            
            if should_search:
                if manual_web_search:
                    status_text = "🔍 Manually searching the web..."
                else:
                    status_text = "🔍 Searching the web for current information..."
            else:
                status_text = "📚 Searching local documents..."
            
            conversation_text = ""
            for msg_dict in history:
                role = msg_dict["role"]
                content = msg_dict["content"]
                if role == "user":
                    conversation_text += f"User: {content}\n"
                else:
                    conversation_text += f"Assistant: {content}\n"

            if should_search:
                # Search the web and add results to vectorstore temporarily
                web_docs = rag.search_and_extract_web_content(user_message, max_results=3)
                if web_docs:
                    # Create a temporary vectorstore with web results
                    temp_vectorstore = rag.create_vectorstore(web_docs)
                    temp_chain = rag.setup_rag_chain(temp_vectorstore)
                    
                    chain_inputs = {
                        "chat_history": conversation_text,
                        "question": user_message
                    }
                    response = rag.query(temp_chain, chain_inputs)
                    if manual_web_search:
                        status_text = f"✅ Manual web search completed - Found {len(web_docs)} web sources"
                    else:
                        status_text = f"✅ Found {len(web_docs)} web sources"
                else:
                    # Fall back to local documents if web search fails
                    chain_inputs = {
                        "chat_history": conversation_text,
                        "question": user_message
                    }
                    response = rag.query(chain, chain_inputs)
                    if manual_web_search:
                        status_text = "⚠️ Manual web search failed, using local documents"
                    else:
                        status_text = "⚠️ Web search failed, using local documents"
            else:
                # Use local documents only
                chain_inputs = {
                    "chat_history": conversation_text,
                    "question": user_message
                }
                response = rag.query(chain, chain_inputs)
                status_text = "✅ Answered using local documents"

            if remove_think_flag:
                response = re.sub(r'<think>.*?</think>', '', response, flags=re.DOTALL).strip()

            history.append({"role": "user", "content": user_message})
            history.append({"role": "assistant", "content": response})
            
            return history, history, "", status_text

        msg.submit(
            fn=chat_with_status,
            inputs=[msg, state, web_search_toggle],
            outputs=[chatbot, state, msg, status]
        )
        
        submit_btn.click(
            fn=chat_with_status,
            inputs=[msg, state, web_search_toggle],
            outputs=[chatbot, state, msg, status]
        )
        
        web_search_btn.click(
            fn=chat_with_status,
            inputs=[msg, state, web_search_toggle, gr.State(True)],
            outputs=[chatbot, state, msg, status]
        )

        def clear_chat():
            return [], [], "", "Ready"

        clear_btn.click(
            fn=clear_chat,
            inputs=[],
            outputs=[chatbot, state, msg, status]
        )

        demo.launch()


if __name__ == "__main__":
    main()
